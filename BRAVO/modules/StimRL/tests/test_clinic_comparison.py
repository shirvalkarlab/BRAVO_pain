"""Host tests for the blinded in-clinic comparison kit and its analysis."""
import json

import numpy as np
import pandas as pd
import pytest

from StimRL import clinic_comparison as CC
from StimRL.data_pipeline import SafetyModel

SAFETY = SafetyModel(ceiling=2.5)


# Independently constructed synthetic settings; no participant record is used.
HOME = {"freq_hz": 70.0, "amp_mA_Left": 1.0, "amp_mA_Right": 1.0,
        "pw_us_Left": 100.0, "pw_us_Right": 100.0, "left_contact": "L C+2-",
        "cathodes": "synthetic contacts"}
TEST = {**HOME, "amp_mA_Left": 1.5, "amp_mA_Right": 1.5}


def test_the_test_setting_is_inside_every_limit_and_differs_from_home():
    CC.check_settings(TEST, HOME, safety=SAFETY)
    assert max(TEST["amp_mA_Left"], TEST["amp_mA_Right"]) <= 2.5


@pytest.mark.parametrize("bad", [{"amp_mA_Left": 4.6}, {"freq_hz": 260.0}, {"amp_mA_Left": 0.0}])
def test_a_setting_past_a_limit_or_off_on_an_active_contact_is_refused(bad):
    with pytest.raises(ValueError):
        CC.check_settings({**TEST, **bad}, HOME, safety=SAFETY)


def test_the_same_setting_twice_is_refused():
    with pytest.raises(ValueError):
        CC.check_settings(HOME, HOME, safety=SAFETY)


def test_every_pair_holds_one_test_and_one_home_step():
    alloc = CC.make_allocation(8)
    assert len(alloc) == 8 and all(sorted(p) == ["C", "R"] for p in alloc)


def test_the_kit_hash_matches_the_allocation_file(tmp_path):
    info = CC.build_kit(tmp_path, n_pairs=4, home_setting=HOME, test_setting=TEST, safety=SAFETY)
    import hashlib
    assert hashlib.sha256((tmp_path / "allocation.json").read_bytes()).hexdigest() == info["sha256"]


def test_the_rater_sheet_names_no_setting_and_no_allocation(tmp_path):
    CC.build_kit(tmp_path, n_pairs=8, home_setting=HOME, test_setting=TEST, safety=SAFETY)
    from openpyxl import load_workbook
    text = " ".join(str(c.value) for ws in load_workbook(tmp_path / "rater_sheet.xlsx").worksheets
                    for row in ws.iter_rows() for c in row if c.value is not None)
    import re
    assert not re.search(r"\b(mA|Hz|us|µs|TEST|HOME)\b", text)
    for word in ("1.5", "1.0", "70.0", "test setting", "home setting", "R/C"):
        assert word not in text
    prog = " ".join(str(c.value) for ws in load_workbook(tmp_path / "programmer_sheet.xlsx").worksheets
                    for row in ws.iter_rows() for c in row if c.value is not None)
    assert "1.5" in prog and "TEST" in prog and "HOME" in prog


def _ratings(alloc, r_minus_c, seed=0):
    rng = np.random.default_rng(seed)
    rows, step = [], 1
    for i, pair in enumerate(alloc):
        base = 6.0 + rng.normal(0, 0.2)            # one level per pair: the difference is exactly r_minus_c
        for lab in pair:
            rows.append({"step": step, "overall": base + (r_minus_c[i] if lab == "R" else 0.0),
                         "left_leg": base, "back": base, "side_effect": "none"})
            step += 1
    return pd.DataFrame(rows)


def test_analysis_finds_a_one_point_drop_in_every_pair_with_the_smallest_possible_p():
    alloc = CC.make_allocation(8)
    res = CC.analyze(_ratings(alloc, [-1.0] * 8), alloc)
    assert res["n_pairs"] == 8 and res["mean_R_minus_C"] == pytest.approx(-1.0, abs=0.15)
    assert res["p_one_sided"] == pytest.approx(1 / 256) and res["passes"] is True


def test_analysis_with_no_difference_does_not_pass():
    alloc = CC.make_allocation(8)
    res = CC.analyze(_ratings(alloc, [0.0] * 8), alloc)
    assert res["passes"] is False and res["p_one_sided"] > 0.05


def test_a_pair_with_a_missing_rating_is_dropped_and_counted():
    alloc = CC.make_allocation(8)
    df = _ratings(alloc, [-1.0] * 8)
    df.loc[0, "overall"] = np.nan
    res = CC.analyze(df, alloc)
    assert res["n_pairs"] == 7 and res["pairs_dropped"] == [1]


def test_kit_requires_reviewed_settings_and_preserves_registered_allocation(tmp_path):
    with pytest.raises(TypeError):
        CC.build_kit(tmp_path)
    assert not tmp_path.joinpath("allocation.json").exists()
    CC.build_kit(tmp_path, n_pairs=2, home_setting=HOME, test_setting=TEST, safety=SAFETY)
    before = tmp_path.joinpath("allocation.json").read_bytes()
    with pytest.raises(ValueError, match="already exists"):
        CC.build_kit(tmp_path, home_setting=HOME, test_setting=TEST, safety=SAFETY)
    assert tmp_path.joinpath("allocation.json").read_bytes() == before


def test_missing_model_prediction_stays_unknown():
    alloc = ["RC"] * 4
    ratings = _ratings(alloc, [-1.] * 4)
    assert CC.analyze(ratings, alloc)["model_predicted"] is None
    assert CC.analyze(ratings, alloc, model_predicted_gain=-0.25)["model_predicted"] == -0.25
