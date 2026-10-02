"""Blinded, randomised in-clinic comparison of the RL model's setting against the home setting.

Approved by the PI on 2026-10-02 ("Yes, please run the small blinded randomized in clinic
comparison!"); test setting chosen by him from the options offered: the RL model's pick.
Protocol: `artifacts/stim_rl_2026-10-02/clinic_comparison_protocol.md`. Nothing here writes to the
device; a clinician programs every step by hand.

    cd BRAVO/modules
    ~/.venvs/bravo-stim-rl/bin/python -m StimRL.clinic_comparison kit       # once, before the visit
    ~/.venvs/bravo-stim-rl/bin/python -m StimRL.clinic_comparison analyze <filled rater sheet>

`kit` writes, to the gitignored `_agent_bridge/_stim_rl_data/clinic_comparison/`:
  allocation.json         the sealed order of each pair (its SHA-256 goes in the protocol)
  programmer_sheet.xlsx   UNBLINDED: which setting to program at each step. Programmer only.
  rater_sheet.xlsx        BLINDED: step numbers and rating boxes, no setting anywhere.
"""
from __future__ import annotations

import hashlib
import json
import os
import secrets
import sys

import numpy as np
import pandas as pd

from . import config as C
from .data_pipeline import DEFAULT_DATA_DIR, SafetyModel

#: The home setting in force since 2026-09-02 19:15 UTC (long-term record, snapshot of 2026-10-02).
#: The programmer confirms it on the tablet at the start of the visit; if it differs, stop and ask.
C_SETTING = {"freq_hz": 55.0, "amp_mA_Left": 3.0, "amp_mA_Right": 2.5, "pw_us_Left": 100.0,
             "pw_us_Right": 150.0, "left_contact": "L C+2-", "cathodes": "Left 2a-2b-2c / Right 1a-1b-1c-2a-2b-2c"}
#: The RL model's pick for L C+2- from today's state (Q-table actor-critic, delta reward, exploit
#: reading): the home setting of 2026-08-12 to 09-02. Same contacts, rate and widths as home.
R_SETTING = {**C_SETTING, "amp_mA_Left": 3.5, "amp_mA_Right": 3.0}
#: The model's own predicted change, R minus C, in pain points (pre-registered; about zero).
MODEL_PREDICTED_GAIN = -0.09

N_PAIRS = 8
STEP_MIN = 3.0          # each step; rated in its last 30 s (the sheets' steps: median 60 s, 75% <= 120 s)
WASHOUT_MIN = 2.0       # back at home after every test step, before the next step
ALPHA = 0.05
OUT_DIR = os.path.join(DEFAULT_DATA_DIR, "clinic_comparison")
SE_LEVELS = ("none", "mild", "mild-persistent", "moderate", "severe")


def _vec(s):
    return (s["freq_hz"], s["amp_mA_Left"], s["amp_mA_Right"], s["pw_us_Left"], s["pw_us_Right"])


def check_settings(r: dict, c: dict) -> None:
    """Refuse a test setting past any limit, with Left current 0 on an active contact, or equal to
    home. Raises ValueError."""
    safety = SafetyModel()
    for name, s in (("test", r), ("home", c)):
        if safety.violates(_vec(s)):
            raise ValueError(f"{name} setting breaks a limit: {s}")
        if s["left_contact"] != "off (Left 0 mA)" and s["amp_mA_Left"] <= 0:
            raise ValueError(f"{name} setting has Left current 0 on an active contact")
    if _vec(r) == _vec(c) and r["left_contact"] == c["left_contact"]:
        raise ValueError("the test setting equals the home setting")


def make_allocation(n_pairs: int = N_PAIRS) -> list:
    """Order within each pair, from the operating system's random source: "RC" or "CR"."""
    return ["RC" if secrets.randbelow(2) else "CR" for _ in range(n_pairs)]


def _steps(alloc):
    out, k = [], 1
    for i, pair in enumerate(alloc, start=1):
        for lab in pair:
            out.append({"step": k, "pair": i, "arm": lab})
            k += 1
    return out


def build_kit(out_dir=OUT_DIR, n_pairs: int = N_PAIRS) -> dict:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    check_settings(R_SETTING, C_SETTING)
    os.makedirs(out_dir, exist_ok=True)
    alloc = make_allocation(n_pairs)
    payload = {"allocation": alloc, "R": R_SETTING, "C": C_SETTING, "n_pairs": n_pairs,
               "salt": secrets.token_hex(16)}
    ap = os.path.join(out_dir, "allocation.json")
    with open(ap, "w") as f:
        json.dump(payload, f, indent=1)
    sha = hashlib.sha256(open(ap, "rb").read()).hexdigest()

    font, bold = Font(name="Arial", size=11), Font(name="Arial", size=11, bold=True)
    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    def sheet(ws, title, lines, header, rows, widths):
        ws["A1"] = title
        ws["A1"].font = Font(name="Arial", size=13, bold=True)
        r = 2
        for line in lines:
            ws.cell(row=r, column=1, value=line).font = font
            r += 1
        r += 1
        for j, h in enumerate(header, start=1):
            c = ws.cell(row=r, column=j, value=h)
            c.font, c.border, c.alignment = bold, box, wrap
            c.fill = PatternFill("solid", start_color="DDEBF7")
        for row in rows:
            r += 1
            for j, v in enumerate(row, start=1):
                c = ws.cell(row=r, column=j, value=v)
                c.font, c.border, c.alignment = font, box, wrap
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[chr(64 + j)].width = w

    steps = _steps(alloc)
    # ---- programmer (unblinded) ----
    wb = Workbook()
    ws = wb.active
    ws.title = "Programmer"
    fmt = lambda s: (f"{s['freq_hz']:.0f} Hz; Left {s['amp_mA_Left']:.1f} mA {s['pw_us_Left']:.0f} us; "
                     f"Right {s['amp_mA_Right']:.1f} mA {s['pw_us_Right']:.0f} us; {s['left_contact']}")
    lines = [
        "UNBLINDED. Programmer only. Do not show this sheet to the patient or the rater.",
        f"HOME = {fmt(C_SETTING)} ({C_SETTING['cathodes']}). Confirm on the tablet before step 1; if it differs, stop.",
        f"TEST = {fmt(R_SETTING)} (same contacts). Only the two currents change (+0.5 mA each side).",
        f"Each step {STEP_MIN:.0f} min; the rater asks for ratings in its last 30 s. After every TEST step, "
        f"{WASHOUT_MIN:.0f} min at HOME before the next step.",
        "At EVERY step change, open the program screen and confirm a value (re-enter the same value when "
        "the step is HOME after HOME), so the patient cannot tell from your actions.",
        "Stop rules: side effect mild-persistent during TEST -> return to HOME, mark the pair, continue; "
        "moderate or severe -> HOME and end the session; patient asks -> end; TEST rated 2+ points worse "
        "than HOME in two pairs in a row -> stop TEST steps.",
        f"Allocation SHA-256 (pre-registered): {sha}",
    ]
    rows = []
    for s in steps:
        setting = R_SETTING if s["arm"] == "R" else C_SETTING
        rows.append([s["step"], s["pair"], "TEST" if s["arm"] == "R" else "HOME",
                     setting["amp_mA_Left"], setting["amp_mA_Right"], "", "",
                     f"then {WASHOUT_MIN:.0f} min HOME" if s["arm"] == "R" else ""])
    sheet(ws, "RCS08 blinded comparison: programmer sheet", lines,
          ["Step", "Pair", "Program", "Left mA", "Right mA", "Time programmed", "Side effect seen / stop rule used", "After"],
          rows, [7, 6, 10, 9, 10, 16, 32, 18])
    wb.save(os.path.join(out_dir, "programmer_sheet.xlsx"))

    # ---- rater (blinded) ----
    wb = Workbook()
    ws = wb.active
    ws.title = "Ratings"
    lines = [
        "Blinded rater sheet. You and the patient are not told what is programmed at any step.",
        "In the last half-minute of each step ask, in this order: overall pain, left leg, back (0 to 10), then side effects.",
        "Side effect: none, mild, mild-persistent, moderate or severe. Tell the programmer at once if moderate or severe.",
        "Write the clock time when you ask. Leave a box empty if the patient does not answer; never guess.",
    ]
    rows = [[s["step"], "", "", "", "", "", ""] for s in steps]
    sheet(ws, "RCS08 blinded comparison: rater sheet", lines,
          ["Step", "Clock time", "Overall pain 0-10", "Left leg 0-10", "Back 0-10", "Side effect", "Comment"],
          rows, [7, 12, 16, 14, 12, 18, 40])
    wb.save(os.path.join(out_dir, "rater_sheet.xlsx"))
    return {"sha256": sha, "out_dir": out_dir, "n_steps": len(steps)}


def _sign_flip_p(d: np.ndarray) -> float:
    """Exact one-sided p for a mean below zero: share of all 2^n sign patterns with a mean as low."""
    n = len(d)
    signs = np.array(np.meshgrid(*[[1.0, -1.0]] * n)).reshape(n, -1).T
    null = (signs * np.abs(d)).mean(axis=1)
    return float(np.mean(null <= d.mean() + 1e-12))


def analyze(ratings: pd.DataFrame, alloc: list, outcome: str = "overall") -> dict:
    """Pair differences R minus C (negative = the test setting hurt less). Pre-registered rule: the
    test setting is better if the one-sided exact sign-flip p is below ALPHA. A pair missing either
    rating is dropped and listed."""
    by_step = ratings.set_index("step")
    diffs, dropped = [], []
    for i, pair in enumerate(alloc, start=1):
        steps = {lab: 2 * (i - 1) + j + 1 for j, lab in enumerate(pair)}
        r = by_step[outcome].get(steps["R"], np.nan)
        c = by_step[outcome].get(steps["C"], np.nan)
        if pd.isna(r) or pd.isna(c):
            dropped.append(i)
        else:
            diffs.append(float(r) - float(c))
    d = np.asarray(diffs)
    out = {"outcome": outcome, "n_pairs": int(len(d)), "pairs_dropped": dropped,
           "differences_R_minus_C": d.tolist()}
    if len(d) < 3:
        out.update(mean_R_minus_C=float("nan"), p_one_sided=float("nan"), passes=False)
        return out
    from scipy.stats import t as tdist
    m, se = float(d.mean()), float(d.std(ddof=1) / np.sqrt(len(d)))
    half = float(tdist.ppf(0.975, len(d) - 1) * se) if se > 0 else 0.0
    out.update(mean_R_minus_C=m, ci95=[m - half, m + half], p_one_sided=_sign_flip_p(d),
               smallest_possible_p=1.0 / 2 ** len(d), model_predicted=MODEL_PREDICTED_GAIN)
    out["passes"] = bool(out["p_one_sided"] < ALPHA)
    return out


def read_rater_sheet(path: str) -> pd.DataFrame:
    raw = pd.read_excel(path, header=None)
    hdr = raw.index[raw.iloc[:, 0].astype(str).str.strip() == "Step"][0]
    df = raw.iloc[hdr + 1:, :7]
    df.columns = ["step", "time", "overall", "left_leg", "back", "side_effect", "comment"]
    df = df[pd.to_numeric(df["step"], errors="coerce").notna()].copy()
    df["step"] = df["step"].astype(int)
    for c in ("overall", "left_leg", "back"):
        df[c] = pd.to_numeric(df[c], errors="coerce")
    return df


def main(argv):
    if not argv or argv[0] not in ("kit", "analyze"):
        raise SystemExit(__doc__)
    if argv[0] == "kit":
        if os.path.exists(os.path.join(OUT_DIR, "allocation.json")):
            raise SystemExit("an allocation already exists; a second one would break the pre-registration")
        print(json.dumps(build_kit(), indent=1))
        return
    alloc = json.load(open(os.path.join(OUT_DIR, "allocation.json")))["allocation"]
    df = read_rater_sheet(argv[1])
    res = {o: analyze(df, alloc, o) for o in ("overall", "left_leg", "back")}
    res["side_effects"] = df["side_effect"].fillna("").astype(str).str.lower().value_counts().to_dict()
    path = os.path.join(OUT_DIR, "result.json")
    json.dump(res, open(path, "w"), indent=1, default=float)
    print(json.dumps(res, indent=1, default=float))


if __name__ == "__main__":
    main(sys.argv[1:])
