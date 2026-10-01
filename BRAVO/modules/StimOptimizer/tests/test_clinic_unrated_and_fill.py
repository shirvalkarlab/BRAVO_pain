"""Unrated clinic steps are kept as EXPOSURE, and their gaps are filled from the visit's own Notes
tab, then from REDCap (the PI, 2026-10-01: "count unrated steps as exposure and look to REDCap to
see if that can help fill in gaps where the pain surveys are not on the STIM tab ... before you
check REDCap, make sure you check the other tabs").

Values, on a constructed workbook laid out like the lab's own: a rated step; a two-row step whose
rating sits on its test row (not "unrated"); an unrated step with a timestamp, filled from a Notes
rating that Excel stored as a date ("8/10" -> 10 August); an unrated step with no timestamp whose
start is inferred from the step before it; and a REDCap survey filed while a step was in force.
"""
import datetime as dt
import os
import tempfile

import numpy as np
import pandas as pd
import pytest

from StimOptimizer import clinic_pain as CP

HEADER = ["Group", "Contacts", "sEEG Contacts", "Amp (mA)", "Rate (Hz)", "PW (µs)", "Threshold",
          "Duration (s)", "Side Effect?", "Timestamp", "Movement/Change point",
          "General Notes / Pt Verbal Notes", "Overall", "Head", "Back", "Left Leg", "Left Foot",
          "Right Leg", "Right Foot"]


def _workbook(path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stim Testing"
    for c, h in enumerate(HEADER, start=1):
        ws.cell(11, c, h)

    def row(r, contacts=None, amp=None, rate=None, pw=None, dur=None, ts=None, overall=None, left_leg=None):
        vals = {2: contacts, 4: amp, 5: rate, 6: pw, 8: dur, 10: ts, 13: overall, 16: left_leg}
        for c, v in vals.items():
            if v is not None:
                ws.cell(r, c, v)
    row(12, "C+2-9-10-", 1.0, 55, 100, 60, dt.time(12, 0, 0), overall=5, left_leg=6)     # rated
    row(13, "C+2-9-10-", 2.0, 55, 100, 60, dt.time(12, 1, 0))                            # ramp row ...
    row(14, dur=60, ts=dt.time(12, 2, 0), overall=4, left_leg=4)                          # ... rated test row
    row(15, "C+1-2-9-10-", 1.6, 55, 100, 120, dt.time(12, 3, 0))                          # unrated, timed
    row(16, "C+1-2-9-10-", 0.0, 55, 100, 60)                                              # unrated, no time
    notes = wb.create_sheet("Notes")
    for c, h in enumerate(["Date", "Time", "Medication", "Activity", "Current verbal pain score",
                           "Head", "BACK", "Left LEG"], start=1):
        notes.cell(1, c, h)
    notes.cell(2, 2, dt.time(12, 4, 30))                     # inside the 12:03 step (120 s)
    notes.cell(2, 5, dt.datetime(2026, 8, 10))               # "8/10" stored by Excel as a date
    notes.cell(2, 8, "7/10")
    notes.cell(3, 2, dt.time(11, 0, 0))                      # before any step: belongs to none
    notes.cell(3, 5, 3)
    wb.save(path)


@pytest.fixture(scope="module")
def parsed():
    d = tempfile.mkdtemp(prefix="clinic_unrated_")
    p = os.path.join(d, "RCS08 Stage 2 - October 2026 Clinic Testing 10_05_26.xlsx")
    _workbook(p)
    return CP.parse_workbook(p, keep_unrated=True)


def test_every_step_is_kept_and_says_where_its_rating_came_from(parsed):
    src = dict(zip(parsed["row_index"], parsed["rating_source"]))
    assert src[12] == "stim_tab" and src[14] == "stim_tab"
    assert src[15] == "notes_tab"
    assert src[16] is None or (isinstance(src[16], float) and np.isnan(src[16]))


def test_a_two_row_steps_ramp_row_is_not_counted_as_an_unrated_step(parsed):
    assert 13 not in set(parsed["row_index"])


def test_a_notes_rating_stored_as_a_date_fills_the_step_in_force_when_it_was_given(parsed):
    r = parsed[parsed["row_index"] == 15].iloc[0]
    assert r["overall"] == 8.0 and r["left_leg"] == 7.0
    assert r["amp_mA_Left"] == 1.6


def test_a_step_with_no_timestamp_starts_where_the_step_before_it_ended(parsed):
    r = parsed[parsed["row_index"] == 16].iloc[0]
    assert r["t_local"] == dt.time(12, 5, 0) and bool(r["t_inferred"]) is True


def test_the_rated_only_view_is_what_every_existing_reader_gets(parsed):
    rated = CP.rated_steps(parsed)
    assert sorted(rated["row_index"]) == [12, 14, 15]


def test_a_redcap_survey_fills_the_step_in_force_when_it_was_filed_and_no_other():
    steps = pd.DataFrame([
        dict(row_index=1, t_utc=pd.Timestamp("2026-01-01 20:00:00", tz="UTC"), duration_s=60.0,
             rating_source=None, overall=None, left_leg=None, back=None),
        dict(row_index=2, t_utc=pd.Timestamp("2026-01-01 20:01:00", tz="UTC"), duration_s=60.0,
             rating_source=None, overall=None, left_leg=None, back=None),
        dict(row_index=3, t_utc=pd.Timestamp("2026-01-01 20:02:00", tz="UTC"), duration_s=60.0,
             rating_source="stim_tab", overall=3.0, left_leg=3.0, back=3.0),
    ])
    reports = pd.DataFrame({"t_utc": [pd.Timestamp("2026-01-01 20:01:40", tz="UTC"),
                                      pd.Timestamp("2026-01-01 20:02:30", tz="UTC"),
                                      pd.Timestamp("2026-01-01 23:00:00", tz="UTC")],
                            "nrs": [6.0, 9.0, 1.0], "left_leg_vas": [70.0, 90.0, 10.0],
                            "back_vas": [50.0, 90.0, 10.0]})
    out = CP.fill_from_redcap(steps, reports)
    r = out.set_index("row_index")
    assert r.loc[2, "rating_source"] == "redcap"
    assert (r.loc[2, "overall"], r.loc[2, "left_leg"], r.loc[2, "back"]) == (6.0, 7.0, 5.0)
    assert pd.isna(r.loc[1, "overall"])                      # its own window had no survey
    assert r.loc[3, "overall"] == 3.0                        # a sheet rating is never replaced


def test_exposure_counts_a_table_with_no_contact_column_as_unrecorded():
    # regression: a steps table from before 2026-10-01 has no contacts_raw; the count must not fail
    d = pd.DataFrame({"amp_mA_Left": [1.0, 0.0], "rating_source": ["stim_tab", None],
                      "visit_date": ["v1", "v1"]})
    out = CP.exposure_by_contact(d)
    labels = {r["left_contact"]: r for r in out}
    assert labels["unrecorded"]["n_steps"] == 1 and labels["off (Left 0 mA)"]["n_rated"] == 0


def test_exposure_by_contact_says_what_each_left_contact_received(parsed):
    labels = {r["left_contact"]: r for r in CP.exposure_by_contact(parsed)}
    assert labels["L C+1-2-"]["n_steps"] == 1 and labels["L C+1-2-"]["amp_max_mA"] == 1.6
    assert labels["L C+1-2-"]["n_rated"] == 1                    # filled from the Notes tab
    assert labels["off (Left 0 mA)"]["n_rated"] == 0


def test_a_rated_step_with_no_timestamp_keeps_no_time_but_still_bounds_the_next_step():
    # regression (2026-10-01): inference wrote times onto 19 RATED RCS08 steps; only unrated
    # steps may take an inferred start, though every step's inferred start bounds the next
    rows = [dict(row_index=1, t_local=dt.time(12, 0), duration_s=60.0, rating_source="stim_tab"),
            dict(row_index=2, t_local=None, duration_s=60.0, rating_source="stim_tab"),
            dict(row_index=3, t_local=None, duration_s=60.0, rating_source=None)]
    CP._infer_missing_starts(rows)
    assert rows[1]["t_local"] is None and rows[1]["t_inferred"] is False
    assert rows[1]["_t_match"] == dt.time(12, 1)
    assert rows[2]["t_local"] == dt.time(12, 2) and rows[2]["t_inferred"] is True


def test_an_unrated_step_with_no_time_at_all_is_a_plan_not_exposure():
    # found live (2026-10-01): the 09_24_26 sheet is the titration card's exported ladder, never
    # filled in -- no timestamp, no rating, an empty Notes tab -- so its 4.5 mA steps were planned,
    # not delivered
    d = pd.DataFrame({
        "amp_mA_Left": [1.0, 4.5, 2.0], "contacts_raw": ["L C+1-2- / R C+1-2-"] * 3,
        "rating_source": [None, None, "stim_tab"], "visit_date": ["v1"] * 3,
        "t_local": [dt.time(12, 0), None, None]})
    labels = {r["left_contact"]: r for r in CP.exposure_by_contact(d)}
    c = labels["L C+1-2-"]
    assert c["n_steps"] == 2 and c["n_rated"] == 1 and c["amp_max_mA"] == 2.0
    assert c["n_planned_only"] == 1
